from __future__ import annotations

from datetime import date
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT_PATH = Path(__file__).with_name("Business-Time-Mileage-Tracker.xlsx")

NAVY = "17324D"
BLUE = "2F6B8A"
TEAL = "2A9D8F"
GOLD = "E9C46A"
CREAM = "F7F4ED"
PALE_BLUE = "EAF2F5"
PALE_TEAL = "E8F5F2"
WHITE = "FFFFFF"
DARK = "24313A"
GRAY = "667782"
LIGHT_GRAY = "D9E2E7"
RED = "C94C4C"

THIN_GRAY = Side(style="thin", color=LIGHT_GRAY)


def style_title(ws, title: str, subtitle: str, end_column: int) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_column)
    title_cell = ws.cell(1, 1, title)
    title_cell.fill = PatternFill("solid", fgColor=NAVY)
    title_cell.font = Font(name="Aptos Display", size=22, bold=True, color=WHITE)
    title_cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 38

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    subtitle_cell = ws.cell(2, 1, subtitle)
    subtitle_cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
    subtitle_cell.font = Font(name="Aptos", size=10, color=GRAY, italic=True)
    subtitle_cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[2].height = 25


def add_section_heading(ws, row: int, title: str, end_column: int) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=end_column)
    cell = ws.cell(row, 1, title)
    cell.fill = PatternFill("solid", fgColor=BLUE)
    cell.font = Font(name="Aptos", size=12, bold=True, color=WHITE)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 24


def make_table(ws, name: str, ref: str, style_name: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name=style_name,
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def configure_log_sheet(
    ws,
    widths: list[float],
    print_area: str,
    tab_color: str,
) -> None:
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"
    ws.sheet_properties.tabColor = tab_color
    ws.print_title_rows = "1:1"
    ws.print_area = print_area
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "Business Time & Mileage Tracker"
    ws.oddFooter.right.text = "Page &P of &N"
    ws.oddFooter.center.size = 8
    ws.oddFooter.right.size = 8

    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + index)].width = width

    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Aptos", bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def add_validations(time_ws, mileage_ws, lists_ws) -> None:
    activity_validation = DataValidation(
        type="list",
        formula1="'Lists'!$A$2:$A$11",
        allow_blank=False,
    )
    activity_validation.promptTitle = "Choose an activity"
    activity_validation.prompt = "Select a common activity or add one on the Lists sheet."
    activity_validation.errorTitle = "Select a listed activity"
    activity_validation.error = "Choose an activity from the dropdown."
    activity_validation.errorStyle = "stop"
    activity_validation.showErrorMessage = True
    activity_validation.showInputMessage = True
    time_ws.add_data_validation(activity_validation)
    activity_validation.add("C2:C1001")

    for ws in (time_ws, mileage_ws):
        date_validation = DataValidation(
            type="date",
            operator="between",
            formula1="DATE(2000,1,1)",
            formula2="DATE(2100,12,31)",
            allow_blank=False,
        )
        date_validation.promptTitle = "Enter a date"
        date_validation.prompt = "Use a date between 2000 and 2100."
        date_validation.errorTitle = "Invalid date"
        date_validation.error = "Enter a valid date between 2000 and 2100."
        date_validation.errorStyle = "stop"
        date_validation.showErrorMessage = True
        date_validation.showInputMessage = True
        ws.add_data_validation(date_validation)
        date_validation.add("A2:A1001")

    hours_validation = DataValidation(
        type="decimal",
        operator="between",
        formula1="0.01",
        formula2="24",
        allow_blank=False,
    )
    hours_validation.promptTitle = "Enter hours"
    hours_validation.prompt = "Enter decimal hours from 0.01 to 24, such as 1.5."
    hours_validation.errorTitle = "Invalid hours"
    hours_validation.error = "Hours must be a number greater than 0 and no more than 24."
    hours_validation.errorStyle = "stop"
    hours_validation.showErrorMessage = True
    hours_validation.showInputMessage = True
    time_ws.add_data_validation(hours_validation)
    hours_validation.add("D2:D1001")

    miles_validation = DataValidation(
        type="decimal",
        operator="between",
        formula1="0.01",
        formula2="10000",
        allow_blank=False,
    )
    miles_validation.promptTitle = "Enter miles"
    miles_validation.prompt = "Enter the business miles driven for this trip."
    miles_validation.errorTitle = "Invalid mileage"
    miles_validation.error = "Miles must be a number greater than 0."
    miles_validation.errorStyle = "stop"
    miles_validation.showErrorMessage = True
    miles_validation.showInputMessage = True
    mileage_ws.add_data_validation(miles_validation)
    miles_validation.add("F2:F1001")

    lists_ws.sheet_state = "hidden"


def build_instructions(ws) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = GOLD
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 24
    ws.column_dimensions["C"].width = 74
    ws.column_dimensions["D"].width = 4
    style_title(
        ws,
        "Business Time & Mileage Tracker",
        "A simple Excel template for rental-property business records",
        4,
    )

    add_section_heading(ws, 4, "Quick start", 4)
    steps = [
        ("1", "Set the reporting year", "Open Summary and enter the year you want to review in cell B3."),
        ("2", "Replace the sample rows", "Open Time Log and Mileage Log, then replace the examples with your own records."),
        ("3", "Add entries", "Type directly in the first blank row under each table. Excel expands the table and formulas automatically."),
        ("4", "Review totals", "Use Summary for yearly totals, monthly detail, and charts."),
        ("5", "Filter or print", "Use the filter arrows in either log. Print settings repeat the header row and fit the log to one page wide."),
    ]
    for row, (number, heading, detail) in enumerate(steps, start=5):
        ws.cell(row, 1, number)
        ws.cell(row, 2, heading)
        ws.cell(row, 3, detail)
        ws.cell(row, 1).fill = PatternFill("solid", fgColor=TEAL)
        ws.cell(row, 1).font = Font(name="Aptos", bold=True, color=WHITE)
        ws.cell(row, 1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(row, 2).font = Font(name="Aptos", bold=True, color=DARK)
        ws.cell(row, 3).font = Font(name="Aptos", color=DARK)
        for column in range(1, 4):
            ws.cell(row, column).border = Border(bottom=THIN_GRAY)
            ws.cell(row, column).alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 38

    add_section_heading(ws, 11, "Tips for clean records", 4)
    tips = [
        ("Use decimal hours", "Record 1 hour 30 minutes as 1.5 hours."),
        ("Keep business purpose specific", "Write a clear reason, such as appliance pickup or tenant inspection."),
        ("Use consistent property names", "Consistent names make filters and future summaries easier."),
        ("Keep source documents", "Retain calendars, receipts, invoices, and other records that support each entry."),
    ]
    for row, (heading, detail) in enumerate(tips, start=12):
        ws.cell(row, 2, heading)
        ws.cell(row, 3, detail)
        ws.cell(row, 2).font = Font(name="Aptos", bold=True, color=BLUE)
        ws.cell(row, 3).font = Font(name="Aptos", color=DARK)
        for column in range(2, 4):
            ws.cell(row, column).border = Border(bottom=THIN_GRAY)
            ws.cell(row, column).alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 34

    add_section_heading(ws, 17, "Start a new year", 4)
    ws.merge_cells("B18:C20")
    ws["B18"] = (
        "Save a copy of this workbook with the year in the file name. Clear the data rows in both logs, "
        "enter the new year on Summary, and keep the prior workbook as your archive. Do not delete the "
        "table header rows."
    )
    ws["B18"].alignment = Alignment(vertical="top", wrap_text=True)
    ws["B18"].font = Font(name="Aptos", color=DARK)
    ws["B18"].fill = PatternFill("solid", fgColor=CREAM)
    ws["B18"].border = Border(left=THIN_GRAY, right=THIN_GRAY, top=THIN_GRAY, bottom=THIN_GRAY)

    add_section_heading(ws, 22, "Customize activity choices", 4)
    ws.merge_cells("B23:C24")
    ws["B23"] = (
        "Unhide the Lists sheet (Home > Format > Hide & Unhide > Unhide Sheet), edit the activity names, "
        "then hide it again. Keep the choices within cells A2:A11 so the dropdown continues to work."
    )
    ws["B23"].alignment = Alignment(vertical="top", wrap_text=True)
    ws["B23"].font = Font(name="Aptos", color=DARK)

    ws.freeze_panes = "A4"
    ws.print_area = "A1:D24"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "Business Time & Mileage Tracker"


def build_summary(ws, year: int) -> None:
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = TEAL
    for column, width in {
        "A": 4,
        "B": 17,
        "C": 16,
        "D": 4,
        "E": 17,
        "F": 16,
        "G": 4,
        "H": 17,
        "I": 16,
    }.items():
        ws.column_dimensions[column].width = width

    style_title(
        ws,
        "Business Activity Summary",
        "Change the reporting year below. Totals and charts update automatically.",
        9,
    )

    ws["A3"] = "Reporting year"
    ws["A3"].font = Font(name="Aptos", bold=True, color=DARK)
    ws["B3"] = year
    ws["B3"].number_format = "0"
    ws["B3"].fill = PatternFill("solid", fgColor=GOLD)
    ws["B3"].font = Font(name="Aptos", size=13, bold=True, color=NAVY)
    ws["B3"].alignment = Alignment(horizontal="center")
    year_validation = DataValidation(
        type="whole",
        operator="between",
        formula1="2000",
        formula2="2100",
        allow_blank=False,
    )
    year_validation.errorTitle = "Invalid year"
    year_validation.error = "Enter a four-digit year between 2000 and 2100."
    year_validation.showErrorMessage = True
    ws.add_data_validation(year_validation)
    year_validation.add(ws["B3"])

    cards = [
        ("B5", "Total hours", '=SUMIFS(TimeLog[Hours],TimeLog[Date],">="&DATE($B$3,1,1),TimeLog[Date],"<"&DATE($B$3+1,1,1))', PALE_TEAL, "0.00"),
        ("D5", "Total miles", '=SUMIFS(MileageLog[Miles],MileageLog[Date],">="&DATE($B$3,1,1),MileageLog[Date],"<"&DATE($B$3+1,1,1))', PALE_BLUE, "0.0"),
        ("F5", "Time entries", '=COUNTIFS(TimeLog[Date],">="&DATE($B$3,1,1),TimeLog[Date],"<"&DATE($B$3+1,1,1))', "F4EBD2", "0"),
        ("H5", "Mileage entries", '=COUNTIFS(MileageLog[Date],">="&DATE($B$3,1,1),MileageLog[Date],"<"&DATE($B$3+1,1,1))', "F4EBD2", "0"),
    ]
    for coordinate, label, formula, fill, number_format in cards:
        column = ws[coordinate].column
        ws.merge_cells(start_row=5, start_column=column, end_row=5, end_column=column + 1)
        ws.merge_cells(start_row=6, start_column=column, end_row=7, end_column=column + 1)
        label_cell = ws.cell(5, column, label)
        value_cell = ws.cell(6, column, formula)
        label_cell.fill = PatternFill("solid", fgColor=NAVY)
        label_cell.font = Font(name="Aptos", bold=True, color=WHITE)
        label_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.fill = PatternFill("solid", fgColor=fill)
        value_cell.font = Font(name="Aptos Display", size=24, bold=True, color=NAVY)
        value_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.number_format = number_format
        for row in range(5, 8):
            for card_column in range(column, column + 2):
                ws.cell(row, card_column).border = Border(
                    left=THIN_GRAY,
                    right=THIN_GRAY,
                    top=THIN_GRAY,
                    bottom=THIN_GRAY,
                )
    ws.row_dimensions[5].height = 22
    ws.row_dimensions[6].height = 31
    ws.row_dimensions[7].height = 18

    ws["B9"] = "Month"
    ws["C9"] = "Hours"
    ws["D9"] = "Miles"
    for cell in ws[9][1:4]:
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Aptos", bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center")

    for row in range(10, 22):
        month_number = row - 9
        ws.cell(row, 2, f"=DATE($B$3,{month_number},1)")
        ws.cell(row, 2).number_format = "mmm"
        ws.cell(
            row,
            3,
            f'=SUMIFS(TimeLog[Hours],TimeLog[Date],">="&B{row},TimeLog[Date],"<"&EDATE(B{row},1))',
        )
        ws.cell(
            row,
            4,
            f'=SUMIFS(MileageLog[Miles],MileageLog[Date],">="&B{row},MileageLog[Date],"<"&EDATE(B{row},1))',
        )
        ws.cell(row, 3).number_format = "0.00"
        ws.cell(row, 4).number_format = "0.0"
        for column in range(2, 5):
            ws.cell(row, column).border = Border(bottom=THIN_GRAY)
            ws.cell(row, column).alignment = Alignment(horizontal="center")

    ws.conditional_formatting.add(
        "C10:C21",
        ColorScaleRule(
            start_type="min",
            start_color=WHITE,
            end_type="max",
            end_color=TEAL,
        ),
    )
    ws.conditional_formatting.add(
        "D10:D21",
        ColorScaleRule(
            start_type="min",
            start_color=WHITE,
            end_type="max",
            end_color=BLUE,
        ),
    )

    hours_chart = BarChart()
    hours_chart.title = "Hours by month"
    hours_chart.style = 10
    hours_chart.y_axis.title = "Hours"
    hours_chart.x_axis.title = "Month"
    hours_chart.height = 7.2
    hours_chart.width = 12.3
    hours_chart.add_data(Reference(ws, min_col=3, min_row=9, max_row=21), titles_from_data=True)
    hours_chart.set_categories(Reference(ws, min_col=2, min_row=10, max_row=21))
    hours_chart.legend = None
    ws.add_chart(hours_chart, "F9")

    mileage_chart = LineChart()
    mileage_chart.title = "Miles by month"
    mileage_chart.style = 13
    mileage_chart.y_axis.title = "Miles"
    mileage_chart.x_axis.title = "Month"
    mileage_chart.height = 7.2
    mileage_chart.width = 12.3
    mileage_chart.add_data(Reference(ws, min_col=4, min_row=9, max_row=21), titles_from_data=True)
    mileage_chart.set_categories(Reference(ws, min_col=2, min_row=10, max_row=21))
    mileage_chart.legend = None
    ws.add_chart(mileage_chart, "F23")

    ws.freeze_panes = "B9"
    ws.print_area = "A1:I36"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.oddFooter.center.text = "Business Time & Mileage Tracker"
    ws.oddFooter.right.text = "Page &P of &N"


def build_time_log(ws, year: int) -> None:
    rows = [
        ["Date", "Property", "Activity", "Hours", "Notes"],
        [date(year, 1, 8), "Maple Street Duplex", "Maintenance", 2.5, "Replaced kitchen faucet"],
        [date(year, 1, 15), "Maple Street Duplex", "Tenant communication", 0.75, "Lease renewal call"],
        [date(year, 2, 3), "Oak Avenue Rental", "Bookkeeping", 1.25, "Reconciled January expenses"],
    ]
    for row in rows:
        ws.append(row)
    configure_log_sheet(ws, [13, 25, 24, 12, 52], "A1:E50", TEAL)
    for cell in ws["A"][1:]:
        cell.number_format = "mmm d, yyyy"
    for cell in ws["D"][1:]:
        cell.number_format = "0.00"
    make_table(ws, "TimeLog", "A1:E4", "TableStyleMedium2")


def build_mileage_log(ws, year: int) -> None:
    rows = [
        ["Date", "Property", "Trip Purpose", "Start Location", "Destination", "Miles", "Notes"],
        [date(year, 1, 8), "Maple Street Duplex", "Maintenance visit", "Home office", "Maple Street Duplex", 18.4, "Faucet repair"],
        [date(year, 1, 20), "Oak Avenue Rental", "Supply pickup", "Home office", "Hardware store", 11.2, "Paint and smoke detectors"],
        [date(year, 2, 3), "Oak Avenue Rental", "Property inspection", "Home office", "Oak Avenue Rental", 24.8, "Quarterly walkthrough"],
    ]
    for row in rows:
        ws.append(row)
    configure_log_sheet(ws, [13, 23, 25, 22, 24, 12, 40], "A1:G50", BLUE)
    for cell in ws["A"][1:]:
        cell.number_format = "mmm d, yyyy"
    for cell in ws["F"][1:]:
        cell.number_format = "0.0"
    make_table(ws, "MileageLog", "A1:G4", "TableStyleMedium4")


def build_lists(ws) -> None:
    ws["A1"] = "Activity"
    activities = [
        "Administration",
        "Bookkeeping",
        "Cleaning",
        "Inspections",
        "Maintenance",
        "Marketing",
        "Property research",
        "Tenant communication",
        "Travel planning",
        "Other",
    ]
    for row, activity in enumerate(activities, start=2):
        ws.cell(row, 1, activity)
    ws.column_dimensions["A"].width = 28


def build_workbook(output_path: Path = OUTPUT_PATH) -> None:
    reporting_year = date.today().year
    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    wb.properties.title = "Business Time and Mileage Tracker"
    wb.properties.subject = "Rental-property business time and mileage record template"
    wb.properties.creator = "Business Time and Mileage Tracker"
    wb.properties.keywords = "rental property, time log, mileage log, business records"

    instructions_ws = wb.create_sheet("Instructions")
    summary_ws = wb.create_sheet("Summary")
    time_ws = wb.create_sheet("Time Log")
    mileage_ws = wb.create_sheet("Mileage Log")
    lists_ws = wb.create_sheet("Lists")

    build_instructions(instructions_ws)
    build_summary(summary_ws, reporting_year)
    build_time_log(time_ws, reporting_year)
    build_mileage_log(mileage_ws, reporting_year)
    build_lists(lists_ws)
    add_validations(time_ws, mileage_ws, lists_ws)

    wb.active = 0
    wb.save(output_path)


def validate_workbook(output_path: Path = OUTPUT_PATH) -> None:
    wb = load_workbook(output_path, data_only=False)
    assert wb.sheetnames == ["Instructions", "Summary", "Time Log", "Mileage Log", "Lists"]
    assert wb["Lists"].sheet_state == "hidden"
    assert "TimeLog" in wb["Time Log"].tables
    assert "MileageLog" in wb["Mileage Log"].tables
    assert wb["Time Log"].tables["TimeLog"].ref == "A1:E4"
    assert wb["Mileage Log"].tables["MileageLog"].ref == "A1:G4"
    assert wb["Time Log"].freeze_panes == "A2"
    assert wb["Mileage Log"].freeze_panes == "A2"
    assert len(wb["Time Log"].data_validations.dataValidation) == 3
    assert len(wb["Mileage Log"].data_validations.dataValidation) == 2
    assert len(wb["Summary"].data_validations.dataValidation) == 1
    assert len(wb["Summary"]._charts) == 2
    assert str(wb["Summary"]["B6"].value).startswith("=SUMIFS(TimeLog[Hours]")
    assert str(wb["Summary"]["D6"].value).startswith("=SUMIFS(MileageLog[Miles]")
    assert str(wb["Summary"]["C10"].value).startswith("=SUMIFS(TimeLog[Hours]")
    assert str(wb["Summary"]["D10"].value).startswith("=SUMIFS(MileageLog[Miles]")

    with ZipFile(output_path) as archive:
        names = set(archive.namelist())
        assert "[Content_Types].xml" in names
        assert "xl/vbaProject.bin" not in names


if __name__ == "__main__":
    build_workbook()
    validate_workbook()
    print(f"Created and validated {OUTPUT_PATH.name}")
